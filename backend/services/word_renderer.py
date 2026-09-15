"""Editable Word export of manual resumes, with compact template typography."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core.builder_schemas import BuilderDraft
from services.inline_formatting import parse_inline_formatting

# Widely available in Word; Liberation Serif provides metric-compatible Linux rendering.
WORD_FONT = "Times New Roman"


def render_builder_word(draft: BuilderDraft, destination: Path) -> None:
    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Inches(8.5), Inches(11)
    page.top_margin = page.bottom_margin = page.left_margin = page.right_margin = Inches(.45)
    usable_width = page.page_width - page.left_margin - page.right_margin
    # Match the existing compact resume design, not Word's default theme.
    for name, size, bold in (("Normal", 9.5, False), ("Title", 22, True),
                             ("Heading 1", 10.5, True), ("Heading 2", 10, True),
                             ("Subtitle", 9.5, False), ("List Bullet", 9.5, False)):
        style = document.styles[name]
        style.font.name = WORD_FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.italic = False
        style.font.color.rgb = RGBColor.from_string("1F1F1F")
        fonts = style.element.get_or_add_rPr().rFonts
        for attribute in list(fonts.attrib):
            if "Theme" in attribute or "theme" in attribute:
                del fonts.attrib[attribute]
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + attribute), WORD_FONT)
        properties = style.element.find(qn("w:pPr"))
        if properties is not None:
            for border in list(properties.findall(qn("w:pBdr"))):
                properties.remove(border)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        # Single spacing is portable to Google Docs. Minimum point-based leading
        # was expanded during import, pushing otherwise one-page resumes to two.
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        style.paragraph_format.widow_control = False
        style.paragraph_format.keep_together = False
        style.paragraph_format.keep_with_next = False
        style.paragraph_format.page_break_before = False
        for spacing in list(style.element.get_or_add_rPr().findall(qn("w:spacing"))):
            style.element.rPr.remove(spacing)
    heading_style = document.styles["Heading 1"]
    heading_style.paragraph_format.space_before = Pt(7)
    heading_style.paragraph_format.space_after = Pt(3)
    heading_style.paragraph_format.keep_with_next = True
    document.styles["Heading 2"].paragraph_format.space_before = Pt(5)
    document.styles["Heading 2"].paragraph_format.keep_with_next = True
    document.styles["Subtitle"].paragraph_format.keep_with_next = True

    # Real Word bullet numbering, with hanging indent for wrapped points.
    numbering = document.part.numbering_part.element
    abstract_id = max([int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))] + [-1]) + 1
    num_id = max([int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))] + [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl"); level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "bullet"), ("lvlText", "\u2022"), ("lvlJc", "left")):
        element = OxmlElement("w:" + tag); element.set(qn("w:val"), value); level.append(element)
    properties = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "240"); indent.set(qn("w:hanging"), "180")
    properties.append(indent); level.append(properties); abstract.append(level); numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    reference = OxmlElement("w:abstractNumId"); reference.set(qn("w:val"), str(abstract_id))
    num.append(reference); numbering.append(num)

    def add_formatted_runs(paragraph, text):
        for segment in parse_inline_formatting(text.strip()):
            run = paragraph.add_run(segment.text)
            run.bold = segment.bold
            run.italic = segment.italic

    def bullet(text):
        p = document.add_paragraph(style="List Bullet")
        add_formatted_runs(p, text)
        p.paragraph_format.left_indent = Pt(12)
        p.paragraph_format.first_line_indent = Pt(-9)
        p.paragraph_format.space_after = Pt(.5)
        p.paragraph_format.keep_together = False
        numbering_properties = OxmlElement("w:numPr")
        for tag, value in (("ilvl", "0"), ("numId", str(num_id))):
            element = OxmlElement("w:" + tag); element.set(qn("w:val"), value)
            numbering_properties.append(element)
        p._p.get_or_add_pPr().append(numbering_properties)

    def heading(title):
        p = document.add_paragraph(title.upper(), style="Heading 1")
        borders = OxmlElement("w:pBdr")
        border = OxmlElement("w:bottom")
        for key, value in (("val", "single"), ("sz", "6"), ("color", "1F1F1F"), ("space", "2")):
            border.set(qn("w:" + key), value)
        borders.append(border); p._p.get_or_add_pPr().append(borders)

    # Plain header avoids Word Title theme residue during Google Docs import.
    name = document.add_paragraph()
    name_run = name.add_run(draft.full_name)
    name_run.font.size = Pt(22)
    name_run.bold = True
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(3)
    contact = document.add_paragraph(" | ".join(filter(None, [draft.phone, draft.email, draft.location])))
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in contact.runs:
        run.font.size = Pt(8.5)
    contact.paragraph_format.space_after = Pt(2)
    mark_properties = OxmlElement("w:rPr")
    mark_size = OxmlElement("w:sz")
    mark_size.set(qn("w:val"), "17")
    mark_properties.append(mark_size)
    contact._p.get_or_add_pPr().append(mark_properties)

    blocks = {"skill_groups": ("Technical Skills", draft.skill_groups),
              "experiences": ("Experience", draft.experiences),
              "projects": ("Projects", draft.projects), "education": ("Education", draft.education)}
    blocks.update({f"custom:{section.id or i}": (section.name, section.entries)
                   for i, section in enumerate(draft.custom_sections)})
    order = list(dict.fromkeys([*draft.section_order, *blocks]))
    for key in order:
        if key not in blocks or not blocks[key][1]:
            continue
        title, entries = blocks[key]
        heading(title)
        for entry in entries:
            if key == "skill_groups":
                p = document.add_paragraph()
                p.add_run(entry.name.strip() + ": ").bold = True
                p.add_run(", ".join(s.strip() for s in entry.skills if s.strip()))
                continue
            p = document.add_paragraph(style="Heading 2")
            p.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(entry.name + (f", {entry.location}" if entry.location else ""))
            if entry.stack and key == "projects":
                stack = p.add_run(" | " + entry.stack); stack.bold = False; stack.italic = True
            if entry.dates:
                date = p.add_run("\t" + entry.dates); date.font.size = Pt(9.6)
            if entry.title:
                document.add_paragraph(entry.title, style="Subtitle")
            if key == "education":
                if entry.coursework.strip():
                    bullet("Relevant coursework: " + entry.coursework)
            else:
                for point in entry.points:
                    if point.strip():
                        bullet(point)
    # Materialize spacing on paragraphs as well as styles. Importers can map
    # named styles to their own defaults; direct spacing preserves our intent.
    for paragraph in document.paragraphs:
        formatting = paragraph.paragraph_format
        inherited = paragraph.style.paragraph_format
        formatting.line_spacing = 1.0
        formatting.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if formatting.space_before is None:
            formatting.space_before = inherited.space_before or Pt(0)
        if formatting.space_after is None:
            formatting.space_after = inherited.space_after or Pt(0)
        formatting.widow_control = False
        formatting.page_break_before = False
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = "Resume"
    document.save(destination)
