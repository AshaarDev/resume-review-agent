"""Resume parsing utilities for extracting text from various file formats."""

import base64
import io
from typing import Dict, Any
from pathlib import Path


class ResumeParser:
    """Parse resumes from PDF, DOCX, images, and text files."""
    
    @staticmethod
    def parse_pdf(file_path: str) -> Dict[str, Any]:
        """Extract text and metadata from PDF resume."""
        try:
            import PyPDF2
            
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                
                text_content = ""
                for page in pdf_reader.pages:
                    text_content += page.extract_text() + "\n"
                
                metadata = {
                    "num_pages": len(pdf_reader.pages),
                    "file_type": "pdf",
                    "file_size": Path(file_path).stat().st_size
                }
                
                return {
                    "text": text_content.strip(),
                    "metadata": metadata,
                    "success": True
                }
        except Exception as e:
            return {
                "text": "",
                "metadata": {},
                "success": False,
                "error": str(e)
            }
    
    @staticmethod
    def parse_docx(file_path: str) -> Dict[str, Any]:
        """Extract text from DOCX resume."""
        try:
            from docx import Document
            
            doc = Document(file_path)
            
            text_content = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)
            
            metadata = {
                "num_paragraphs": len(doc.paragraphs),
                "file_type": "docx",
                "file_size": Path(file_path).stat().st_size
            }
            
            return {
                "text": "\n".join(text_content),
                "metadata": metadata,
                "success": True
            }
        except Exception as e:
            return {
                "text": "",
                "metadata": {},
                "success": False,
                "error": str(e)
            }
    
    @staticmethod
    def parse_image(file_path: str) -> Dict[str, Any]:
        """Extract text from image resume using OCR."""
        try:
            from PIL import Image
            import pytesseract
            
            image = Image.open(file_path)
            text_content = pytesseract.image_to_string(image)
            
            metadata = {
                "image_size": image.size,
                "image_mode": image.mode,
                "file_type": "image",
                "file_size": Path(file_path).stat().st_size
            }
            
            return {
                "text": text_content.strip(),
                "metadata": metadata,
                "success": True
            }
        except Exception as e:
            return {
                "text": "",
                "metadata": {},
                "success": False,
                "error": str(e)
            }
    
    @staticmethod
    def parse_from_base64(base64_data: str, file_type: str) -> Dict[str, Any]:
        """Parse resume from base64 encoded data."""
        try:
            file_data = base64.b64decode(base64_data)
            
            if file_type.lower() == 'pdf':
                import PyPDF2
                pdf_file = io.BytesIO(file_data)
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                
                text_content = ""
                for page in pdf_reader.pages:
                    text_content += page.extract_text() + "\n"
                
                return {
                    "text": text_content.strip(),
                    "metadata": {"file_type": "pdf", "num_pages": len(pdf_reader.pages)},
                    "success": True
                }
            
            elif file_type.lower() in ['jpg', 'jpeg', 'png', 'gif', 'bmp']:
                from PIL import Image
                import pytesseract
                
                image = Image.open(io.BytesIO(file_data))
                text_content = pytesseract.image_to_string(image)
                
                return {
                    "text": text_content.strip(),
                    "metadata": {"file_type": "image", "image_size": image.size},
                    "success": True
                }
            
            elif file_type.lower() == 'docx':
                from docx import Document
                
                doc = Document(io.BytesIO(file_data))
                text_content = []
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip():
                        text_content.append(paragraph.text)
                
                return {
                    "text": "\n".join(text_content),
                    "metadata": {"file_type": "docx"},
                    "success": True
                }
            
            else:
                return {
                    "text": "",
                    "metadata": {},
                    "success": False,
                    "error": f"Unsupported file type: {file_type}"
                }
                
        except Exception as e:
            return {
                "text": "",
                "metadata": {},
                "success": False,
                "error": str(e)
            }
    
    @staticmethod
    def parse_resume(file_path: str) -> Dict[str, Any]:
        """Parse resume from any supported format (auto-detect)."""
        file_ext = Path(file_path).suffix.lower()
        
        if file_ext == '.pdf':
            return ResumeParser.parse_pdf(file_path)
        elif file_ext in ['.docx', '.doc']:
            return ResumeParser.parse_docx(file_path)
        elif file_ext in ['.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tiff']:
            return ResumeParser.parse_image(file_path)
        elif file_ext == '.txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
            return {
                "text": text,
                "metadata": {"file_type": "text"},
                "success": True
            }
        else:
            return {
                "text": "",
                "metadata": {},
                "success": False,
                "error": f"Unsupported file format: {file_ext}"
            }
