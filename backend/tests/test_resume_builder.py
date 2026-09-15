"""Manual builder template safety, PNG preview, and durable exports."""
import fitz
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from services.word_renderer import render_builder_word
from fastapi.testclient import TestClient
from main import app
from core.builder_schemas import BuilderDraft
from routes import builder
from services.resume_builder import render_builder_latex
from services.latex_compiler import LatexCompilationResult
from core.creator_schemas import CompilationStatus

client = TestClient(app)


def test_word_export_is_editable_ordered_and_has_real_bullets(tmp_path):
    data = intake()
    data["custom_sections"] = [{"id": "awards", "name": "Awards", "entries": [
        {"name": "Merit award", "points": ["Won recognition for engineering."]}]}]
    data["section_order"] = ["custom:awards", "education", "experiences", "projects", "skill_groups"]
    path = tmp_path / "resume.docx"
    render_builder_word(BuilderDraft.model_validate(data), path)
    document = Document(path)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Ada & Grace" in text and "Saved 40%" in text
    assert text.index("AWARDS") < text.index("EDUCATION") < text.index("EXPERIENCE")
    assert "React, Python" in text
    bullets = [p for p in document.paragraphs if p.style.name == "List Bullet"]
    assert bullets and all(p._p.pPr.numPr is not None for p in bullets)
    assert document.sections[0].left_margin.inches == .45
    assert document.styles["Normal"].font.size.pt == 9.5
    assert document.styles["Normal"].font.name == "Times New Roman"
    assert document.styles["Heading 2"].font.name == "Times New Roman"
    assert document.styles["Normal"].paragraph_format.line_spacing == 1.0
    assert document.styles["Normal"].paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE
    assert all(p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE for p in document.paragraphs)
    assert all(p.paragraph_format.space_before is not None and p.paragraph_format.space_after is not None
               for p in document.paragraphs)
    assert document.paragraphs[0].style.name != "Title"
    assert document.styles["Normal"].paragraph_format.widow_control is False
    assert all(p.paragraph_format.keep_together is False for p in bullets)
    assert document.styles["Heading 1"].paragraph_format.keep_with_next is True


def intake():
    return {"full_name": "Ada & Grace", "experiences": [
        {"name": "Company", "title": "Engineer", "points": ["Saved 40% by using caching."]}],
        "education": [{"name": "University", "title": "Computer Science", "coursework": "Algorithms, Databases"}],
        "projects": [{"name": "Dashboard", "stack": "React, Python", "points": ["Built an analytics dashboard."]}],
        "skill_groups": [{"name": "Languages", "skills": ["Python", "SQL"]}]}


def fake_compile(path):
    pdf_path = path.with_suffix(".pdf")
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), "User-authored resume content")
        pdf.save(pdf_path)
    return LatexCompilationResult(CompilationStatus.COMPILED, pdf_path=pdf_path)


def test_builder_reuses_template_and_escapes_user_text():
    data = intake()
    data["experiences"][0]["points"].append(r"\input{secret}")
    tex = render_builder_latex(BuilderDraft.model_validate(data))
    assert r"Ada \& Grace" in tex
    assert r"Saved 40\%" in tex
    assert r"\input{secret}" not in tex
    assert "Relevant coursework: Algorithms, Databases" in tex
    assert "Dashboard" in tex and r"\textit{React, Python}" in tex
    assert "Languages" in tex


def test_builder_renders_safe_inline_bold_and_italic(tmp_path):
    data = intake()
    data["experiences"][0]["points"] = [
        "Saved **40%** by using *request caching*."
    ]
    data["education"][0]["coursework"] = (
        "**Algorithms**, *Database Systems*"
    )
    draft = BuilderDraft.model_validate(data)

    tex = render_builder_latex(draft)
    assert r"Saved \textbf{40\%} by using \textit{request caching}." in tex
    assert r"\textbf{Algorithms}, \textit{Database Systems}" in tex

    path = tmp_path / "formatted-resume.docx"
    render_builder_word(draft, path)
    document = Document(path)
    bullet = next(
        p for p in document.paragraphs
        if p.style.name == "List Bullet" and "Saved 40%" in p.text
    )
    assert "**" not in bullet.text and "*" not in bullet.text
    assert any(run.text == "40%" and run.bold for run in bullet.runs)
    assert any(
        run.text == "request caching" and run.italic
        for run in bullet.runs
    )


def test_builder_repairs_nested_and_legacy_editor_emphasis(tmp_path):
    data = intake()
    data["experiences"][0]["points"] = [
        "****Built a platform.**** Led **a *team *of 3**."
    ]
    draft = BuilderDraft.model_validate(data)

    tex = render_builder_latex(draft)
    assert "****" not in tex
    assert r"\textbf{Built a platform.}" in tex
    assert r"\textbf{a }" in tex
    assert r"\textbf{\textit{team }}" in tex
    assert r"\textbf{of 3}" in tex

    path = tmp_path / "legacy-formatted-resume.docx"
    render_builder_word(draft, path)
    document = Document(path)
    bullet = next(
        p for p in document.paragraphs
        if p.style.name == "List Bullet" and "Built a platform" in p.text
    )
    assert "*" not in bullet.text
    assert any(run.text == "Built a platform." and run.bold for run in bullet.runs)
    assert any(run.text == "team " and run.bold and run.italic for run in bullet.runs)


def test_builder_preview_returns_png_and_cleans_workspace(monkeypatch):
    paths = []
    def compiler(path):
        paths.append(path.parent)
        return fake_compile(path)
    monkeypatch.setattr(builder, "compile_latex", compiler)
    response = client.post("/api/resume-builder/preview", json=intake())
    assert response.status_code == 200
    assert response.content.startswith(b"\x89PNG")
    assert response.headers["x-page-count"] == "1"
    assert not paths[0].exists()


def test_builder_source_pdf_handoff_is_temporary(monkeypatch):
    paths = []
    def compiler(path):
        paths.append(path.parent)
        return fake_compile(path)
    monkeypatch.setattr(builder, "compile_latex", compiler)
    response = client.post("/api/resume-builder/source-pdf", json=intake())
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["cache-control"] == "no-store"
    assert response.content.startswith(b"%PDF")
    assert not paths[0].exists()


def test_builder_export_saves_before_workspace_cleanup(monkeypatch):
    monkeypatch.setattr(builder, "compile_latex", fake_compile)
    def save(tex, pdf, word):
        assert tex.exists() and pdf.exists() and word.exists()
        return "a" * 32
    monkeypatch.setattr(builder, "save_resume_artifacts", save)
    response = client.post("/api/resume-builder/export", json=intake())
    assert response.status_code == 200
    assert response.json()["pdf_download_url"].endswith("/resume.pdf")
    assert response.json()["docx_download_url"].endswith("/resume.docx")


def test_builder_rejects_empty_export_and_oversize_points():
    assert client.post("/api/resume-builder/export", json={}).status_code == 422
    data = intake()
    data["experiences"][0]["points"] = ["x" * 701]
    assert client.post("/api/resume-builder/preview", json=data).status_code == 422


def test_builder_reports_missing_pdf_compiler(monkeypatch):
    monkeypatch.setattr(builder, "compile_latex", lambda _: LatexCompilationResult(
        CompilationStatus.SOURCE_ONLY, error_message="PDF compiler unavailable."))
    response = client.post("/api/resume-builder/preview", json=intake())
    assert response.status_code == 503


def test_builder_rejects_missing_page(monkeypatch):
    monkeypatch.setattr(builder, "compile_latex", fake_compile)
    assert client.post("/api/resume-builder/preview?page=2", json=intake()).status_code == 404


def test_custom_only_sections_render_and_escape():
    draft = BuilderDraft.model_validate({"full_name": "Ada", "custom_sections": [
        {"name": "Leadership & Awards", "entries": [
            {"name": "Club", "dates": "2025", "points": [r"Won 50% recognition; \input{secret}"]}]}]})
    tex = render_builder_latex(draft)
    assert r"\section{LEADERSHIP \& AWARDS}" in tex
    assert r"Won 50\%" in tex
    assert r"\input{secret}" not in tex
    assert tex.index("LEADERSHIP") < tex.index(r"\end{document}")


def test_builder_preserves_section_entry_and_point_order():
    data = intake()
    data["custom_sections"] = [{"id": "awards", "name": "Awards", "entries": [
        {"name": "Winner", "points": ["First custom point", "Second custom point"]}]}]
    data["section_order"] = ["custom:awards", "education", "projects", "experiences", "skill_groups"]
    data["experiences"].insert(0, {"name": "Newest employer", "title": "Lead", "points": [
        "First achievement", "Second achievement"]})
    tex = render_builder_latex(BuilderDraft.model_validate(data))
    headings = [r"\section{AWARDS}", r"\section{EDUCATION}", r"\section{PROJECTS}",
                r"\section{EXPERIENCE}", r"\section{TECHNICAL SKILLS}"]
    positions = [tex.index(heading) for heading in headings]
    assert positions == sorted(positions)
    assert tex.index("Newest employer") < tex.index("Company")
    assert tex.index("First achievement") < tex.index("Second achievement")
    data["section_order"] = ["education", "education", r"\input{secret}", "missing"]
    tex = render_builder_latex(BuilderDraft.model_validate(data))
    assert tex.count(r"\section{EDUCATION}") == 1
    assert r"\input{secret}" not in tex
    assert all(heading in tex for heading in headings)


def test_custom_sections_export_and_bounds(monkeypatch):
    monkeypatch.setattr(builder, "compile_latex", fake_compile)
    monkeypatch.setattr(builder, "save_resume_artifacts", lambda *_: "a" * 32)
    data = {"full_name": "Ada", "custom_sections": [{"name": "Awards", "entries": [
        {"name": "Scholarship", "points": ["Earned a merit scholarship."]}]}]}
    assert client.post("/api/resume-builder/export", json=data).status_code == 200
    assert client.post("/api/resume-builder/preview", json=data).status_code == 200
    data["custom_sections"][0]["entries"] = []
    assert client.post("/api/resume-builder/export", json=data).status_code == 422
    data["custom_sections"] *= 9
    assert client.post("/api/resume-builder/preview", json=data).status_code == 422
